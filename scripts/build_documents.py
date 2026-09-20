"""Build the academic documents from Markdown using python-docx and ReportLab.

Authoring dependencies: python-docx, reportlab, Pillow.
The documents remain reproducible from the readable Markdown sources.
"""
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image as PDFImage, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
FIGURES = DOCS / 'figures'
FIGURES.mkdir(exist_ok=True)
INK = '#24364b'


def font(size, bold=False):
    candidates = [Path('C:/Windows/Fonts') / ('arialbd.ttf' if bold else 'arial.ttf'),
                  Path('/usr/share/fonts/truetype/dejavu') / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default(size=size)


def diagram(name, title, boxes, arrows, size=(1600, 760)):
    image = Image.new('RGB', size, 'white')
    draw = ImageDraw.Draw(image)
    draw.text((45, 25), title, fill=INK, font=font(34, True))
    for start, end, label in arrows:
        draw.line([start, end], fill='#72859a', width=4)
        x, y = end
        if abs(end[0]-start[0]) > abs(end[1]-start[1]):
            direction = 1 if end[0] > start[0] else -1
            points = [(x,y), (x-14*direction,y-8), (x-14*direction,y+8)]
        else:
            direction = 1 if end[1] > start[1] else -1
            points = [(x,y), (x-8,y-14*direction), (x+8,y-14*direction)]
        draw.polygon(points, fill='#72859a')
        if label:
            midpoint = ((start[0]+end[0])/2, (start[1]+end[1])/2)
            draw.text((midpoint[0]+10, midpoint[1]+8), label, fill='#53667b', font=font(20))
    for rect, heading, text in boxes:
        draw.rounded_rectangle(rect, radius=12, fill='#f2f5f8', outline='#98aabc', width=2)
        x1,y1,x2,y2 = rect
        draw.text((x1+20,y1+18), heading, fill=INK, font=font(27, True))
        draw.multiline_text((x1+20,y1+62), text, fill='#43586f', font=font(23), spacing=10)
    path = FIGURES / f'{name}.png'
    image.save(path)


def figures():
    diagram('architecture', 'System architecture', [
        ((40,130,340,290),'Browser','Forms and file list\nUntrusted input'),
        ((470,110,1000,310),'Waitress and Flask','Authentication and CSRF\nValidation and normalization\nOwner checks and audit'),
        ((1150,110,1550,310),'Optional ClamAV','Normalized bytes\nBounded local process\nReject errors or detection'),
        ((130,470,530,680),'SQLite','Users and sessions\nFile metadata and events'),
        ((610,470,1010,680),'Private vault','Random file identifiers\nFernet encrypted blobs'),
        ((1090,470,1490,680),'Private keys','Session signing\nEncryption and audit HMAC'),
    ], [((340,210),(470,210),''),((1000,210),(1150,210),''),((610,310),(330,470),''),
        ((810,310),(810,470),''),((980,310),(1290,470),'')])
    diagram('upload-flow', 'Upload validation and storage sequence', [
        ((40,115,500,260),'1  Authenticate and bound','Session and CSRF\nRequest and file size limits'),
        ((570,115,1030,260),'2  Inspect declared type','Filename and extension\nMIME consistency'),
        ((1100,115,1560,260),'3  Validate content','Magic bytes and decoder\nPixel limit or UTF-8 rules'),
        ((1100,405,1560,560),'4  Normalize and scan','Image reconstruction\nOptional antivirus'),
        ((570,405,1030,560),'5  Enforce account quota','50 MiB and 200 files\nSerialized quota check'),
        ((40,405,500,560),'6  Encrypt and commit','Random ID, hash, ciphertext\nMetadata and audit event'),
        ((370,650,1230,735),'Any failed check rejects the upload',''),
    ], [((500,190),(570,190),''),((1030,190),(1100,190),''),((1330,260),(1330,405),''),
        ((1100,480),(1030,480),''),((570,480),(500,480),'')], size=(1600,800))
    diagram('database', 'Database relationships', [
        ((40,185,490,475),'USERS','id PK, username UNIQUE\npassword_hash\ntotp_secret, totp_pending\ntotp_last, created_at'),
        ((770,90,1540,270),'FILES','id PK, owner_id FK, name, size, mime\nsha256, scan, created_at'),
        ((770,320,1540,475),'SESSIONS','token_hash PK, user_id FK\nexpires_at'),
        ((770,525,1540,715),'EVENTS','id PK, user_id FK (nullable)\naction, detail, created_at, chain_mac'),
    ], [((490,250),(770,180),'1 to many'),((490,340),(770,395),'1 to many'),
        ((490,440),(770,615),'1 to many')])


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows=[]
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [part.strip() for part in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', item) for item in row):
                    rows.append(row)
                i += 1
            yield 'table', rows
            continue
        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            yield 'image', image.groups()
        elif line.startswith('# '):
            yield 'title', line[2:]
        elif line.startswith('### '):
            yield 'h2', line[4:]
        elif line.startswith('## '):
            yield 'h1', line[3:]
        elif line.startswith('- '):
            yield 'bullet', line[2:]
        else:
            yield 'paragraph', line
        i += 1


def clean(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('**','').replace('`','')


def make_docx(source, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ['Normal','Body Text','List Bullet']:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(11)
        doc.styles[name].paragraph_format.space_after = Pt(7)
        doc.styles[name].paragraph_format.line_spacing = 1.15
    for name,size in [('Title',27),('Heading 1',17),('Heading 2',12)]:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0,0,0)
        doc.styles[name].paragraph_format.space_before = Pt(15)
        doc.styles[name].paragraph_format.space_after = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Secure File Upload System  |  ').font.size = Pt(9)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    first_section = True
    for kind,value in blocks(source.read_text(encoding='utf-8')):
        if kind=='title':
            doc.add_paragraph(clean(value), 'Title')
        elif kind=='h1':
            if first_section:
                doc.add_page_break()
                doc.add_heading('Contents', 1)
                # An editable field. Word updates it using References > Update Table.
                p = doc.add_paragraph()
                field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
                p._p.append(field)
                doc.add_paragraph('Update the contents field in Word after final edits.', style='Caption')
                doc.add_page_break()
                first_section=False
            doc.add_heading(clean(value),1)
        elif kind=='h2': doc.add_heading(clean(value),2)
        elif kind=='table':
            table=doc.add_table(rows=1, cols=len(value[0]))
            table.style='Light Shading Accent 1'
            for index,cell in enumerate(value[0]):
                table.rows[0].cells[index].text=clean(cell)
            repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in value[1:]:
                cells=table.add_row().cells
                for index,text in enumerate(row): cells[index].text=clean(text)
            for row in table.rows:
                prevent=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(prevent)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(5)
                        for run in p.runs: run.font.size=Pt(10)
            doc.add_paragraph()
        elif kind=='image':
            title,path=value
            doc.add_picture(str(DOCS/path), width=Inches(6.45))
            doc.add_paragraph(title,'Caption')
        else:
            doc.add_paragraph(clean(value), 'List Bullet' if kind=='bullet' else 'Normal')
    doc.core_properties.title='Secure File Upload System'
    doc.core_properties.subject=source.stem.replace('_',' ')
    doc.core_properties.author=''
    doc.save(output)


def make_pdf(source, output):
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='BodyProject', fontName='Helvetica',fontSize=10.5,leading=15,spaceAfter=8,textColor=colors.HexColor('#25303c')))
    styles.add(ParagraphStyle(name='CoverProject',fontName='Helvetica-Bold',fontSize=29,leading=35,spaceAfter=25,textColor=colors.black))
    styles.add(ParagraphStyle(name='H1Project',fontName='Helvetica-Bold',fontSize=16,leading=21,spaceBefore=17,spaceAfter=9,keepWithNext=True))
    styles.add(ParagraphStyle(name='H2Project',fontName='Helvetica-Bold',fontSize=11.5,leading=16,spaceBefore=11,spaceAfter=7,keepWithNext=True))
    styles.add(ParagraphStyle(name='CellProject',fontName='Helvetica',fontSize=8.8,leading=12,spaceAfter=0))
    styles.add(ParagraphStyle(name='CaptionProject',fontName='Helvetica-Oblique',fontSize=9,leading=12,spaceAfter=12,alignment=TA_CENTER,textColor=colors.HexColor('#526374')))
    styles.add(ParagraphStyle(name='ContentsProject',fontName='Helvetica',fontSize=10.5,leading=17,spaceAfter=3))
    def paragraph(text,style='BodyProject'):
        value=escape(clean(text)).replace('\u2013','-').replace('\u2014','-').replace('\u2011','-')
        # Permit long reference URLs to wrap without shrinking body type.
        value=re.sub(r'https?://\S+',lambda m: m.group().replace('/','/<wbr/>').replace('_','_<wbr/>'),value)
        return Paragraph(value,styles[style])
    source_blocks=list(blocks(source.read_text(encoding='utf-8')))
    story=[]; first=True
    for kind,value in source_blocks:
        if kind=='title':
            story.extend([Spacer(1,60),paragraph(value,'CoverProject')])
        elif kind=='h1':
            if first:
                story.extend([Spacer(1,24),paragraph('Academic project documentation'),PageBreak(),paragraph('Contents','H1Project')])
                contents = [[paragraph(item,'ContentsProject')] for item_kind,item in source_blocks if item_kind=='h1']
                contents_table = Table(contents, colWidths=[A4[0]-102], hAlign='LEFT')
                contents_table.setStyle(TableStyle([('LEFTPADDING',(0,0),(-1,-1),0),
                                                   ('TOPPADDING',(0,0),(-1,-1),2),
                                                   ('BOTTOMPADDING',(0,0),(-1,-1),2)]))
                story.append(contents_table)
                story.append(PageBreak()); first=False
            story.append(paragraph(value,'H1Project'))
        elif kind=='h2': story.append(paragraph(value,'H2Project'))
        elif kind=='table':
            columns=len(value[0]); width=(A4[0]-102)/columns
            cells=[[paragraph(cell,'CellProject') for cell in row] for row in value]
            table=Table(cells,colWidths=[width]*columns,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf0f5')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),
                ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
                ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
                ('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#d7dfe7')),
            ]))
            story.extend([table,Spacer(1,11)])
        elif kind=='image':
            title,path=value; image=Image.open(DOCS/path)
            width=A4[0]-102; height=width*image.height/image.width
            story.append(KeepTogether([PDFImage(str(DOCS/path),width=width,height=height),paragraph(title,'CaptionProject')]))
        elif kind=='bullet': story.append(paragraph('• '+value))
        else: story.append(paragraph(value))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#617183'))
        canvas.drawString(51,30,'Secure File Upload System')
        canvas.drawRightString(A4[0]-51,30,str(doc.page)); canvas.restoreState()
    SimpleDocTemplate(str(output),pagesize=A4,leftMargin=51,rightMargin=51,topMargin=48,bottomMargin=48,
                      title='Secure File Upload System',author='').build(story,onFirstPage=footer,onLaterPages=footer)


if __name__=='__main__':
    figures()
    for stem in ['SYNOPSIS','PROJECT_REPORT']:
        source=DOCS/f'{stem}.md'
        make_docx(source,DOCS/f'{stem}.docx')
        make_pdf(source,DOCS/f'{stem}.pdf')
        print(f'Created {stem}.docx and {stem}.pdf')
