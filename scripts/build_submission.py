"""Build print-ready PDFs and an editable presentation from project sources.

Authoring dependencies: reportlab, python-pptx, python-docx and Pillow.
Run with the document-authoring Python environment, not the web server.
"""
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                               Spacer, PageBreak, Table, TableStyle, Image as PDFImage, KeepTogether)
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from build_documents import blocks, clean, figures, diagram

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
fonts = Path('C:/Windows/Fonts')
if (fonts/'times.ttf').exists():
    pdfmetrics.registerFont(TTFont('Academic', str(fonts/'times.ttf')))
    pdfmetrics.registerFont(TTFont('AcademicBold', str(fonts/'timesbd.ttf')))
else:
    pdfmetrics.registerFont(pdfmetrics.Font('Academic', 'Times-Roman', 'WinAnsiEncoding'))
    pdfmetrics.registerFont(pdfmetrics.Font('AcademicBold', 'Times-Bold', 'WinAnsiEncoding'))


def roman(number):
    result = ''
    for value, symbol in [(100,'c'), (90,'xc'), (50,'l'), (40,'xl'), (10,'x'), (9,'ix'), (5,'v'), (4,'iv'), (1,'i')]:
        while number >= value:
            result += symbol
            number -= value
    return result


class AcademicPDF(BaseDocTemplate):
    def beforeDocument(self):
        self.body_start = None

    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and getattr(flowable, 'toc_entry', False):
            title = flowable.getPlainText()
            if title.startswith('Chapter 1'):
                self.body_start = self.page
            page = str(self.page-self.body_start+1) if self.body_start else roman(self.page-1)
            key = 'heading-' + str(self.page) + '-' + str(len(title))
            self.canv.bookmarkPage(key)
            self.notify('TOCEntry', (0, title, page, key))


def make_pdf(stem):
    body = ParagraphStyle('body', fontName='Academic', fontSize=12, leading=18,
                          alignment=TA_JUSTIFY, spaceAfter=9)
    styles = {
        'body': body,
        'h1': ParagraphStyle('h1', parent=body, fontName='AcademicBold', fontSize=16, leading=22, alignment=TA_CENTER, spaceAfter=18, keepWithNext=True),
        'h2': ParagraphStyle('h2', parent=body, fontName='AcademicBold', fontSize=14, leading=20, spaceBefore=12, keepWithNext=True, alignment=0),
        'caption': ParagraphStyle('caption', parent=body, fontName='AcademicBold', fontSize=10, leading=15, alignment=TA_CENTER, keepWithNext=False),
        'cell': ParagraphStyle('cell', parent=body, fontSize=10, leading=15, alignment=0, spaceAfter=0),
        'front': ParagraphStyle('front', parent=body, fontSize=14, leading=28),
    }
    def p(text, style='body'):
        text=escape(clean(text))
        text=re.sub(r'https?://\S+', lambda m:m.group().replace('/','/<wbr/>').replace('_','_<wbr/>'), text)
        return Paragraph(text, styles[style])
    source=list(blocks((DOCS/f'{stem}.md').read_text(encoding='utf-8-sig')))
    story=[]
    report=stem=='PROJECT_REPORT'
    toc_added=False
    current=''
    for kind,value in source:
        if kind=='title':
            story.extend([Spacer(1,80),p(value,'h1'),Spacer(1,30)])
        elif kind=='h1':
            story.append(PageBreak())
            if (not toc_added) and (not report or value=='Abstract'):
                story.append(p('Table of Contents','h1'))
                toc=TableOfContents()
                toc.levelStyles=[ParagraphStyle('toc',fontName='Academic',fontSize=12,leading=18,spaceBefore=5)]
                story.extend([toc,PageBreak()])
                if report:
                    for label, entries in [('List of Figures',[v[0] for k,v in source if k=='image']),
                                           ('List of Tables',[v for k,v in source if k=='paragraph' and v.startswith('Table ')])]:
                        heading=p(label,'h1');heading.toc_entry=True;story.append(heading)
                        story.extend(p(entry) for entry in entries)
                        story.append(PageBreak())
                toc_added=True
            heading=p(value,'h1')
            heading.toc_entry=value!='Bonafide Certificate'
            story.append(heading);current=value
        elif kind=='h2':
            story.append(p(value,'h2'))
        elif kind=='table':
            count=len(value[0]); width=A4[0]-112
            cells=[[p(cell,'cell') for cell in row] for row in value]
            table=Table(cells,colWidths=[width/count]*count,repeatRows=1)
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#d9d9d9')),
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
                ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
            story.extend([table,Spacer(1,12)])
        elif kind=='image':
            label,path=value
            with Image.open(DOCS/path) as im:
                scale=min((A4[0]-112)/im.width,390/im.height)
                pic=PDFImage(str(DOCS/path),im.width*scale,im.height*scale)
            story.append(KeepTogether([pic,Spacer(1,6),p(label,'caption')]))
        elif kind=='paragraph' and value.startswith('Table '):
            caption=p(value,'caption');caption.keepWithNext=True;story.append(caption)
        else:
            style='front' if current in ('Bonafide Certificate','Abstract') else 'body'
            story.append(p(('• ' if kind=='bullet' else '')+value,style))
    doc=AcademicPDF(str(DOCS/f'{stem}.pdf'),pagesize=A4,leftMargin=56,rightMargin=56,topMargin=55,bottomMargin=55,title=stem.replace('_',' '),author='')
    def footer(c,d):
        if d.page==1:return
        c.setFont('Academic',10)
        page=str(d.page-d.body_start+1) if d.body_start else roman(d.page-1)
        c.drawCentredString(A4[0]/2,30,page)
    doc.addPageTemplates(PageTemplate(id='main',frames=[Frame(56,55,A4[0]-112,A4[1]-110,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPageEnd=footer))
    doc.multiBuild(story)
    print('Created',stem+'.pdf')


def make_slides():
    suite=ET.parse(DOCS/'test-results.xml').getroot().find('testsuite')
    total=suite.attrib['tests'];failed=suite.attrib['failures'];errors=suite.attrib['errors']
    slides=[
        ('Secure File Upload System','SecureVault | Cybersecurity project', ['[Student name]  /  [Enrollment number]','[Institution]  /  [Guide]','Validation, document scanning and controlled sharing'],None,'Introduce the practical problem and identify your own contribution. Replace the student placeholders before presenting.'),
        ('The problem','Files cross a trust boundary', ['Extensions and MIME declarations can be deceptive.','Images can retain unnecessary location or camera metadata.','Public file directories and reusable links can expose uploads.','A single check cannot establish that a file is safe.'],None,'Explain a renamed file example without demonstrating attacks on external systems.'),
        ('The solution','A private workspace with explicit gates',['Validate content and enforce resource limits.','Reconstruct images and optionally scan document photos.','Encrypt accepted bytes and enforce ownership.','Grant recipients one download through an expiring code.'],'screenshots/vault.png','Show that the workspace contains actual uploaded files. The screenshot uses synthetic demo data.'),
        ('Scope and boundaries','TXT, PNG and JPEG',['5 MiB per file; 16 megapixels per image.','50 MiB and 200 files per account.','No PDF, Office, SVG, archive or executable uploads.','Validation is not a malware-free guarantee.'],None,'Use the narrow format scope as a deliberate engineering decision.'),
        ('Technology stack','Small enough to inspect',['Flask + Jinja + Waitress: web application.','SQLite: transactional metadata and session state.','Pillow + OpenCV: normalization and document processing.','cryptography + PyOTP: encryption and authenticator codes.','pytest + Bandit + pip-audit: verification tooling.'],None,'Distinguish framework, database, image-processing and verification responsibilities.'),
        ('System architecture','The server is a trusted component',[],'figures/architecture.png','Trace the browser request into validation, private storage and metadata. Keys are never public assets.'),
        ('Source structure','Where each concept is implemented',['app/__init__.py: routes, schema and security pipeline','app/document_scanner.py: page detection and cleanup','app/templates + app/static: interface and interactions','tests/: security, sharing and scanner regression cases','scripts/: startup and artifact generation','docs/: synopsis, report, evidence and presentation'],None,'Open these files during the viva if asked where a feature is implemented.'),
        ('Upload pipeline','Independent checks before storage',[],'figures/upload-flow.png','Explain extension, MIME, magic bytes and full decoding as different checks. All enforcement happens on the server.'),
        ('OpenCV document scanner','Photo to readable page',['Validated image input; processing capped at a 2000-pixel side.','Canny edges and convex four-corner contour selection.','Perspective correction, then grayscale CLAHE enhancement.','No detected page: enhance the full frame without cropping.','This is document processing, not OCR or antivirus.'],'screenshots/upload.png','Upload samples/document-photo.png with the checkbox enabled. The synthetic page is suitable for a repeatable demonstration.'),
        ('Storage and authentication','Protect bytes and account access',['Fernet: AES-128-CBC with HMAC-SHA256.','scrypt password hashes; SHA-256 file fingerprints.','Revocable server sessions with 30-minute absolute expiry.','Optional TOTP with encrypted secrets and replay checks.','The server holds keys: encryption at rest, not end-to-end.'],None,'Explain why hashing is not encryption and why losing the key prevents recovery.'),
        ('Single-use sharing','Deliberate recipient access',['Owner creates a random 96-bit code.','Only a hash is stored; code expires after 15 minutes.','Recipient submits code; integrity and expiry are checked.','Atomic consumption permits at most one redemption.','Owner can revoke access before the code is used.'],None,'Possession of the code grants access. An interrupted download may require a new code.'),
        ('Database relationships','Ownership is explicit',[],'figures/database.png','Explain users, sessions, files, events and shares. Foreign keys connect records; transactions protect competing writes.'),
        ('Secure SDLC','Security requirements become test cases',['Requirements: actors, supported formats and threats.','Design: compare local, separate API and managed-cloud options.','Implementation: bounded processing and explicit access checks.','Verification: positive, negative and concurrent test cases.','Deployment and maintenance: HTTPS, backups and dependency review.'],None,'No invented six-week schedule is included. Explain the feedback loop from a failing test to a fix.'),
        ('Research and uniqueness','Integration with visible evidence',['Eian et al. (2020): practical secure-development guidance.','Kudriavtseva and Gadyatskaya (2022): secure SDLC practices.','Neef and Oudeh (2024): varied upload-scanner evaluation.','Contribution: integrated controls, transparent UI and reproducible tests.','No claim of novel cryptography or a new malware detector.'],None,'The report contains a deliberately small three-paper descriptive corpus. Do not present it as a comprehensive bibliometric study.'),
        ('Validation results','Measured evidence, stated limits',[f'{total} tests in the recorded suite; {failed} failures; {errors} errors.','Cases include metadata removal, ownership and code concurrency.','Scanner tests include skewed-page and no-boundary fixtures.','Antivirus subprocess outcomes are mocked in tests.','See test-results.xml and TESTING.md for dates and scan results.'],None,'Read the current evidence rather than memorizing a stale count. This run is not proof that vulnerabilities are absent.'),
        ('Live demonstration','A short reproducible walkthrough',['Sign in and upload samples/demo-notes.txt.','Scan samples/document-photo.png with OpenCV enabled.','Preview and download the processed image.','Create a share code and redeem it once.','Try reusing the code; inspect activity and security settings.'],None,'Prepare two browser sessions. Use synthetic files. Do not expose authenticator setup secrets during a recorded demonstration.'),
        ('Limitations and next steps','What the project does not claim',['ClamAV requires separate operator installation.','Audit suffix deletion needs external checkpoints to detect.','No recovery codes or managed key rotation.','Docker files supplied; local container execution unverified.','Future: passkeys, worker queues and object storage.'],None,'Be candid about operational boundaries. These are future work, not buttons pretending to work.'),
        ('References','Primary sources and reproducible artifacts',['OWASP File Upload and Password Storage Cheat Sheets','NIST SP 800-218 Secure Software Development Framework','OpenCV geometric transformations and CLAHE documentation','Flask, Pillow, cryptography and PyOTP documentation','Full URLs and research records: docs/PROJECT_REPORT.md'],None,'Open the report references for the exact URLs. Invite questions about the design tradeoffs.'),
    ]
    slides = [slides[0],
        ('Problem and solution','Controlled file exchange',['Risk: deceptive types, exposed metadata and uncontrolled downloads.','Solution: validate, normalize, encrypt and enforce ownership.','Recipients receive one download through an expiring code.','Scope: TXT, PNG and JPEG; optional document-photo scanning.'],None,'Explain why several independent checks are required.'),
        ('Technology and file structure','Responsibilities are visible in source',['Flask / SQLite / Waitress: app/__init__.py and run.py','Pillow / OpenCV: validation and app/document_scanner.py','Jinja / CSS / JavaScript: app/templates and app/static','Fernet / scrypt / PyOTP: storage, passwords and MFA','pytest: tests/ | CI: .github/ | startup/builders: scripts/','Report, synopsis, diagrams and evidence: docs/'],None,'Database entities are users, sessions, files, shares and events. The full ER diagram is in the report.'),
        slides[5],
        ('End-to-end workflow','From account to recipient',['Sign in -> choose image/text -> validate type and size.','Normalize -> optional document scan -> optional antivirus.','Check quota -> encrypt -> save metadata and audit event.','Owner previews/downloads or creates a share code.','Recipient redeems once -> record use -> deny code reuse.'],None,'The report expands this workflow into upload and database diagrams. Rejection stops persistence.'),
        slides[8],slides[9],slides[10],
        ('SDLC and research','Requirements become reproducible checks',['Requirements -> threat model -> compare designs -> implement.','Verify positive, negative and concurrent behavior; fix defects.','Deploy with protected keys, HTTPS and dependency review.','Research: Eian (2020), Kudriavtseva (2022), Neef (2024).','Uniqueness: integrated controls, visible state and inspectable tests.'],None,'The three-paper corpus is descriptive, not comprehensive. Research detection rates are not project results.'),
        slides[14],
        ('Live demo and limitations','Show the working path',['Upload demo-notes.txt; scan document-photo.png.','Preview the processed image; create and redeem a code.','Attempt code reuse; inspect activity and account security.','ClamAV requires installation; no OCR or recovery codes.','Future: external audit checkpoints, passkeys and worker queues.'],None,'Use the samples folder. Docker execution and public production hosting are not claimed. Keep authenticator secrets out of screenshots.'),
        slides[17]]
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    pdf=canvas.Canvas(str(DOCS/'PRESENTATION.pdf'),pagesize=(960,540))
    notes=['# Presentation speaker notes','']
    for index,(title,subtitle,bullets,picture,note) in enumerate(slides,1):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F8F5ED')
        pdf.setFillColor(colors.HexColor('#F8F5ED'));pdf.rect(0,0,960,540,fill=1,stroke=0)
        def text(x,y,w,h,value,size=20,bold=False):
            box=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72))
            tf=box.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
            run=tf.paragraphs[0].add_run();run.text=value;run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string('25232E')
            para=Paragraph(escape(value),ParagraphStyle('slide',fontName='Helvetica-Bold' if bold else 'Helvetica',fontSize=size,leading=size*1.25,textColor=colors.HexColor('#25232E')))
            _,height=para.wrap(w,h);para.drawOn(pdf,x,540-y-height)
        text(45,25,850,20,f'SECUREVAULT  /  {index:02d}',11,True)
        text(45,68,870,62,title,34,True)
        text(45,125,860,35,subtitle,17)
        if picture:
            path=DOCS/picture
            with Image.open(path) as im:
                available=(445,300) if bullets else (860,330)
                scale=min(available[0]/im.width,available[1]/im.height)
                w,h=im.width*scale,im.height*scale
            x=490 if bullets else 45+(860-w)/2;y=175
            slide.shapes.add_picture(str(path), Inches(x/72),Inches(y/72),width=Inches(w/72),height=Inches(h/72))
            pdf.drawImage(str(path),x,540-y-h,w,h)
        width=410 if picture else 850
        for n,line in enumerate(bullets):text(45,185+n*55,width,52,line,17 if picture else 20)
        text(45,512,860,18,'Secure File Upload System  |  Academic demonstration',10)
        slide.notes_slide.notes_text_frame.text=note
        notes.extend([f'## Slide {index} {title}',note,''])
        pdf.showPage()
    deck.save(DOCS/'PRESENTATION.pptx');pdf.save()
    (DOCS/'PRESENTATION_NOTES.md').write_text('\n'.join(notes),encoding='utf-8')
    print('Created presentation: 12 slides, editable PPTX and PDF')


if __name__=='__main__':
    figures()
    diagram('database','Logical data relationships',[
        ((530,95,1040,270),'Users','Identity and password hash\nEncrypted TOTP state'),
        ((40,385,490,620),'Sessions','User ID and token hash\nAbsolute expiry'),
        ((570,385,1020,620),'Files','Owner, MIME, size and SHA-256\nRandom blob ID and AV result'),
        ((1090,385,1550,620),'Events','User ID and event fields\nPrior MAC and HMAC'),
        ((570,735,1130,920),'Shares','Owner and file ID\nCode hash, expiry and used time'),
    ],[((650,270),(260,385),'1 : many'),((780,270),(780,385),'1 : many'),((940,270),(1300,385),'1 : many'),((800,620),(800,735),'1 : many')],size=(1600,980))
    for stem in ('SYNOPSIS','PROJECT_REPORT'):make_pdf(stem)
    make_slides()
